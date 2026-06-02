from rest_framework import viewsets, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import AllowAny
from django.db import models
from .models import Factor, Caracteristica, Indicador, Pregunta, Respuesta, Participacion
from .serializers import FactorSerializer, CaracteristicaSerializer, IndicadorSerializer, PreguntaSerializer
from usuarios.models import Estudiante, Docente, Egresado, Administrativo, Directivo, Empleador
from django.contrib.auth import authenticate
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from django.http import HttpResponse


# 1. ViewSets de la estructura del CNA
class FactorViewSet(viewsets.ModelViewSet):
    queryset = Factor.objects.all()
    serializer_class = FactorSerializer
    permission_classes = [AllowAny]

class CaracteristicaViewSet(viewsets.ModelViewSet):
    queryset = Caracteristica.objects.all()
    serializer_class = CaracteristicaSerializer
    permission_classes = [AllowAny]

class IndicadorViewSet(viewsets.ModelViewSet):
    queryset = Indicador.objects.all()
    serializer_class = IndicadorSerializer
    permission_classes = [AllowAny]

class PreguntaViewSet(viewsets.ModelViewSet):
    queryset = Pregunta.objects.all()
    serializer_class = PreguntaSerializer
    permission_classes = [AllowAny]

# 2. Vista para el Login de la Encuesta
class ValidarAccesoEncuestaView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        doc = request.data.get('documento')
        pwd = request.data.get('contrasena')

        if not doc or not pwd or str(doc).strip() != str(pwd).strip():
            return Response({'error': 'El usuario y la contraseña no coinciden.'}, status=status.HTTP_401_UNAUTHORIZED)

        estamentos = {
            'Estudiantes': Estudiante,
            'Docentes': Docente,
            'Egresados': Egresado,
            'Administrativos': Administrativo,
            'Directivos': Directivo,
            'Empleadores': Empleador
        }

        for nombre_estamento, modelo in estamentos.items():
            persona = modelo.objects.filter(dni=doc).first()
            if persona:
                return Response({
                    'mensaje': 'Acceso permitido',
                    'estamento': nombre_estamento,
                    'documento': doc
                }, status=status.HTTP_200_OK)

        return Response({'error': 'El documento no se encuentra registrado en el programa.'}, status=status.HTTP_404_NOT_FOUND)

# 3. Vista para entregar y recibir el Cuestionario
class CuestionarioView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, estamento):
        preguntas = Pregunta.objects.filter(
            models.Q(estamento_dirigido=estamento) | models.Q(estamento_dirigido='Todos')
        )
        serializer = PreguntaSerializer(preguntas, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        estamento = request.data.get('estamento')
        documento = request.data.get('documento') # Por si luego deseas auditar quién respondió
        respuestas = request.data.get('respuestas', [])

        if not respuestas:
            return Response({'error': 'No se enviaron respuestas'}, status=status.HTTP_400_BAD_REQUEST)

        for item in respuestas:
            Respuesta.objects.create(
                pregunta_id=item['pregunta_id'],
                valoracion=item['valoracion'],
                estamento=estamento
            )
            if documento:
                Participacion.objects.get_or_create(documento=documento, defaults={'estamento': estamento})
                
                return Response({'mensaje': '¡Encuesta guardada con éxito!'}, status=status.HTTP_201_CREATED)
            
            
class GraficaParticipacionView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        # Mapeamos los 6 estamentos con sus respectivos modelos
        estamentos = {
            'Estudiantes': Estudiante,
            'Docentes': Docente,
            'Egresados': Egresado,
            'Administrativos': Administrativo,
            'Directivos': Directivo,
            'Empleadores': Empleador
        }

        etiquetas = []
        votaron = []
        faltan = []

        # Calculamos las estadísticas dinámicamente para cada grupo
        for nombre, modelo in estamentos.items():
            total = modelo.objects.count()
            participaron = Participacion.objects.filter(estamento=nombre).count()
            faltantes = max(0, total - participaron)

            etiquetas.append(nombre)
            votaron.append(participaron)
            faltan.append(faltantes)

        datos_grafica = {
            'etiquetas': etiquetas,
            'votaron': votaron,
            'faltan': faltan
        }
        
        return Response(datos_grafica, status=status.HTTP_200_OK)
            
class LoginAdminView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        # Capturamos los datos, sin importar si Angular los envía en inglés o en español
        usuario = request.data.get('username') or request.data.get('usuario')
        contrasena = request.data.get('password') or request.data.get('contrasena')

        # El sistema nativo de Django verifica la contraseña encriptada
        user = authenticate(username=usuario, password=contrasena)

        if user is not None:
            if user.is_staff or user.is_superuser:
                # Si es administrador, le abrimos la puerta y le damos un token de acceso
                return Response({
                    'mensaje': 'Acceso concedido',
                    'token': 'token-admin-autorizado',
                    'usuario': user.username
                }, status=status.HTTP_200_OK)
            else:
                return Response({'error': 'No tienes permisos de administrador para entrar al Dashboard.'}, status=status.HTTP_403_FORBIDDEN)
        else:
            return Response({'error': 'Credenciales incorrectas.'}, status=status.HTTP_401_UNAUTHORIZED)
        
class DetalleParticipacionView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, estamento):
        estamentos = {
            'Estudiantes': Estudiante,
            'Docentes': Docente,
            'Egresados': Egresado,
            'Administrativos': Administrativo,
            'Directivos': Directivo,
            'Empleadores': Empleador
        }

        if estamento not in estamentos:
            return Response({'error': 'Estamento no válido'}, status=status.HTTP_400_BAD_REQUEST)

        modelo = estamentos[estamento]
        # Obtenemos todos los usuarios registrados en ese estamento
        usuarios = modelo.objects.all()
        
        # Obtenemos solo los DNIs de los que ya votaron (usamos set para búsqueda ultrarrápida)
        votantes = set(Participacion.objects.filter(estamento=estamento).values_list('documento', flat=True))

        datos = []
        for u in usuarios:
            datos.append({
                'dni': u.dni,
                'nombre': getattr(u, 'nombre', 'Sin nombre registrado'), # Usa 'nombre' si existe
                'estado': 'Completado' if u.dni in votantes else 'Pendiente'
            })
            
        return Response(datos, status=status.HTTP_200_OK)
class GenerarInformeWordView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        document = Document()

        # Configuración APA 7 (Márgenes y Fuente)
        for section in document.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        style = document.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 2.0

        # --- PORTADA ---
        document.add_paragraph('\n\n\n')
        titulo = document.add_paragraph('Informe de Autoevaluación con Fines de Acreditación')
        titulo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        titulo.runs[0].bold = True
        document.add_paragraph('\n\n')
        info_institucional = document.add_paragraph(
            'Comité de Autoevaluación\n'
            'Programa de Ingeniería en Producción Acuícola\n'
            'Universidad de Nariño\n'
            'San Juan de Pasto\n'
            '2026'
        )
        info_institucional.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_page_break()

        # --- CAPÍTULO 1: RESUMEN EJECUTIVO ---
        h1 = document.add_paragraph('Resumen Ejecutivo')
        h1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        h1.runs[0].bold = True

        document.add_paragraph(
            'El presente documento consolida los resultados preliminares del proceso de autoevaluación '
            'del programa académico. A continuación, se detalla el nivel de participación y cobertura '
            'de los distintos estamentos consultados durante el levantamiento de información.'
        )

        # --- CAPÍTULO 2: ESTADÍSTICAS DE PARTICIPACIÓN ---
        h2 = document.add_paragraph('\nParticipación de la Comunidad Académica')
        h2.runs[0].bold = True

        # Definimos los estamentos a consultar
        estamentos_dict = {
            'Estudiantes': Estudiante,
            'Docentes': Docente,
            'Egresados': Egresado,
            'Administrativos': Administrativo,
            'Directivos': Directivo,
            'Empleadores': Empleador
        }

        # Crear tabla en Word (1 fila de encabezado, 4 columnas)
        tabla = document.add_table(rows=1, cols=4)
        tabla.style = 'Table Grid' # Estilo de cuadrícula clásico
        
        # Llenar encabezados
        encabezados = tabla.rows[0].cells
        encabezados[0].text = 'Estamento'
        encabezados[1].text = 'Población Registrada'
        encabezados[2].text = 'Encuestas Completadas'
        encabezados[3].text = 'Porcentaje (%)'

        # Poner en negrita los encabezados
        for cell in encabezados:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Calcular datos y llenar filas
        total_poblacion = 0
        total_participacion = 0

        for nombre, modelo in estamentos_dict.items():
            poblacion = modelo.objects.count()
            participaron = Participacion.objects.filter(estamento=nombre).count()
            
            porcentaje = 0
            if poblacion > 0:
                porcentaje = (participaron / poblacion) * 100

            total_poblacion += poblacion
            total_participacion += participaron

            # Añadir fila por cada estamento
            fila = tabla.add_row().cells
            fila[0].text = nombre
            fila[1].text = str(poblacion)
            fila[2].text = str(participaron)
            fila[3].text = f"{porcentaje:.1f}%"

        # Añadir fila de Totales al final
        fila_total = tabla.add_row().cells
        fila_total[0].text = 'TOTAL'
        fila_total[1].text = str(total_poblacion)
        fila_total[2].text = str(total_participacion)
        
        porcentaje_global = 0
        if total_poblacion > 0:
            porcentaje_global = (total_participacion / total_poblacion) * 100
        fila_total[3].text = f"{porcentaje_global:.1f}%"

        # Poner en negrita la fila de totales
        for cell in fila_total:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        # --- CAPÍTULO 3: FACTORES Y CARACTERÍSTICAS (CNA) ---
        document.add_page_break()
        h3 = document.add_paragraph('Estructura del Modelo de Autoevaluación (CNA)')
        h3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        h3.runs[0].bold = True

        document.add_paragraph(
            'El proceso de autoevaluación del programa de Ingeniería en Producción Acuícola '
            'se fundamenta en los lineamientos del Consejo Nacional de Acreditación (CNA). '
            'A continuación, se detalla la estructura de factores y características evaluadas:'
        )

        # Consultamos todos los factores de la base de datos
        factores = Factor.objects.all()
        
        for factor in factores:
            # Creamos el párrafo del Factor y lo ponemos en negrita
            p_factor = document.add_paragraph()
            num_factor = getattr(factor, 'numero', factor.id)
            desc_factor = getattr(factor, 'descripcion', getattr(factor, 'nombre', 'Sin descripción'))
            p_factor.add_run(f'Factor {num_factor}: {desc_factor}').bold = True
            
            # Consultamos las características vinculadas específicamente a este factor
            caracteristicas = Caracteristica.objects.filter(factor=factor)
            
            for carac in caracteristicas:
                num_carac = getattr(carac, 'numero', carac.id)
                desc_carac = getattr(carac, 'descripcion', getattr(carac, 'nombre', 'Sin descripción'))
                
                # Agregamos la característica con formato de viñeta nativo de Word
                p_carac = document.add_paragraph(f'Característica {num_carac}: {desc_carac}')
                p_carac.style = 'List Bullet' 
        # --- CAPÍTULO 3: FACTORES, CARACTERÍSTICAS E INDICADORES (CNA) ---
        document.add_page_break()
        h3 = document.add_paragraph('Estructura del Modelo de Autoevaluación (CNA)')
        h3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        h3.runs[0].bold = True

        document.add_paragraph(
            'El proceso de autoevaluación del programa de Ingeniería en Producción Acuícola '
            'se fundamenta en los lineamientos del Consejo Nacional de Acreditación (CNA). '
            'A continuación, se detalla la estructura jerárquica de factores, características e indicadores evaluados:'
        )

        # Consultamos todos los factores
        factores = Factor.objects.all()
        
        for factor in factores:
            # 1. Párrafo del Factor (Negrita)
            p_factor = document.add_paragraph()
            num_factor = getattr(factor, 'numero', factor.id)
            desc_factor = getattr(factor, 'descripcion', getattr(factor, 'nombre', 'Sin descripción'))
            p_factor.add_run(f'\nFactor {num_factor}: {desc_factor}').bold = True
            
            # Consultamos las características de este factor
            caracteristicas = Caracteristica.objects.filter(factor=factor)
            
            for carac in caracteristicas:
                num_carac = getattr(carac, 'numero', carac.id)
                desc_carac = getattr(carac, 'descripcion', getattr(carac, 'nombre', 'Sin descripción'))
                
                # 2. Párrafo de la Característica (Viñeta principal)
                p_carac = document.add_paragraph(f'Característica {num_carac}: {desc_carac}')
                p_carac.style = 'List Bullet' 

                # Consultamos los indicadores vinculados específicamente a esta característica
                indicadores = Indicador.objects.filter(caracteristica=carac)
                
                for ind in indicadores:
                    num_ind = getattr(ind, 'numero', ind.id)
                    desc_ind = getattr(ind, 'descripcion', getattr(ind, 'nombre', 'Sin descripción'))
                    
                    # 3. Párrafo del Indicador (Viñeta secundaria indentada)
                    p_ind = document.add_paragraph(f'Indicador {num_ind}: {desc_ind}')
                    p_ind.style = 'List Bullet'
                    # Lo empujamos media pulgada a la derecha para crear jerarquía visual
                    p_ind.paragraph_format.left_indent = Inches(0.5)

        

        # --- GENERAR ARCHIVO ---
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = 'attachment; filename="Informe_Consolidado_IPA.docx"'
        
        
        document.save(response)
        return response